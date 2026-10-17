"""Generate the editable 15-slide review deck: docs/ThoraxAI_Presentation.pptx

    python scripts/make_presentation.py

Slide 13 (results) is deliberately a placeholder: fill it only with numbers from
results/<model>/test_metrics.json after running training and evaluation.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import config  # noqa: E402

OUT = ROOT / "docs" / "ThoraxAI_Presentation.pptx"
ICON = ROOT / "assets" / "icon.png"

# Palette: deep clinical teal dominates, coral is the single sharp accent.
TEAL = RGBColor(0x14, 0x4E, 0x63)
TEAL_MID = RGBColor(0x1F, 0x6F, 0x8B)
TEAL_TINT = RGBColor(0xE6, 0xF0, 0xF3)
CORAL = RGBColor(0xE0, 0x7A, 0x5F)
CORAL_TINT = RGBColor(0xFB, 0xEA, 0xE4)
INK = RGBColor(0x1D, 0x29, 0x39)
MUTED = RGBColor(0x5B, 0x67, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD_FONT, BODY_FONT = "Cambria", "Calibri"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)


def _run_style(run, size, color, bold=False, italic=False, font=BODY_FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False, font=BODY_FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, name=None):
    """Add a text box. ``content`` is a string or list of strings (one paragraph each)."""
    box = slide.shapes.add_textbox(x, y, w, h)
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, para in enumerate([content] if isinstance(content, str) else content):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        _run_style(p.add_run(), size, color, bold, italic, font)
        p.runs[0].text = para
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.28)))
        pPr.set("indent", str(-Inches(0.28)))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        if isinstance(item, tuple):  # (bold lead, rest)
            r = p.add_run()
            r.text = item[0]
            _run_style(r, size, color, bold=True)
            r2 = p.add_run()
            r2.text = item[1]
            _run_style(r2, size, color)
        else:
            r = p.add_run()
            r.text = item
            _run_style(r, size, color)
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def box_text(shape, content, size=14, color=INK, bold=False, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.1)
    for i, line in enumerate([content] if isinstance(content, str) else content):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        _run_style(r, size if i == 0 else size - 2, color, bold=bold and i == 0)


def badge(slide, x, y, number, fill=TEAL_MID):
    c = rect(slide, x, y, Inches(0.5), Inches(0.5), fill, shape=MSO_SHAPE.OVAL)
    box_text(c, str(number), 14, WHITE, bold=True)
    return c


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def new_slide(prs, title, dark=False, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = TEAL if dark else WHITE
    if title:
        text(slide, MARGIN, Inches(0.45), W - 2 * MARGIN, Inches(0.85), title, size=34, bold=True,
             color=WHITE if dark else TEAL, font=HEAD_FONT, name="Title")
    slide.shapes.add_picture(str(ICON), W - Inches(0.95), H - Inches(0.85), Inches(0.42), Inches(0.42))
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def card(slide, x, y, w, h, heading, body, fill=TEAL_TINT, head_color=TEAL, number=None):
    rect(slide, x, y, w, h, fill)
    inset = Inches(0.25)
    top = y + Inches(0.2)
    if number is not None:
        badge(slide, x + inset, top, number)
        top += Inches(0.62)
    text(slide, x + inset, top, w - 2 * inset, Inches(0.5), heading, 17, head_color, bold=True)
    text(slide, x + inset, top + Inches(0.48), w - 2 * inset, h - (top - y) - Inches(0.6), body, 14, INK)


def build() -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 Title ---------------------------------------------------------------
    s = new_slide(prs, None, dark=True, notes="Introduce the team and the one-line aim: multi-label chest X-ray "
                  "finding classification with explanations, built as an educational prototype.")
    s.shapes.add_picture(str(ICON), MARGIN, Inches(1.2), Inches(1.3), Inches(1.3))
    text(s, MARGIN, Inches(2.75), Inches(11), Inches(1.1), "ThoraxAI", 54, WHITE, bold=True, font=HEAD_FONT)
    text(s, MARGIN, Inches(3.75), Inches(11), Inches(1.0),
         "Multi-Label Chest X-ray Disease Classification Using Deep Learning and Explainable AI", 22, TEAL_TINT)
    text(s, MARGIN, Inches(5.0), Inches(7), Inches(1.4),
         ["<Student 1>  ·  <Student 2>  ·  <Student 3>  ·  <Student 4>",
          "Guide: <Guide name>   |   BE CSE (AI & ML)   |   <College>, <Year>"], 16, WHITE)
    tag = rect(s, Inches(8.6), Inches(5.15), Inches(4.1), Inches(0.6), CORAL)
    box_text(tag, "Educational prototype · not for clinical use", 14, WHITE, bold=True)

    # 2 Introduction ----------------------------------------------------------
    s = new_slide(prs, "Introduction", notes="Context: chest radiographs are common; deep CNNs can learn visual "
                  "patterns; public datasets make this feasible. Our aim is a transparent, honest pipeline.")
    stats = [("14", "thoracic findings scored independently"), ("1", "frontal X-ray in, 14 scores out"),
             ("0", "simulated predictions or metrics shown")]
    for i, (big, small) in enumerate(stats):
        x = MARGIN + i * Inches(4.1)
        rect(s, x, Inches(1.6), Inches(3.8), Inches(1.9), TEAL_TINT)
        text(s, x + Inches(0.3), Inches(1.7), Inches(3.2), Inches(1.0), big, 54, CORAL, bold=True, font=HEAD_FONT)
        text(s, x + Inches(0.3), Inches(2.75), Inches(3.3), Inches(0.7), small, 15, INK)
    bullets(s, MARGIN, Inches(3.9), Inches(12), Inches(3), [
        ("Chest radiography ", "is one of the most frequently performed imaging examinations; reading volume is high."),
        ("Deep CNNs ", "learn hierarchical image features and transfer well from natural images (ImageNet)."),
        ("ChestX-ray14 ", "(NIH) provides ~112k labelled frontal images, enabling multi-finding models."),
        ("ThoraxAI ", "builds the full pipeline — validation, training, evaluation, explanation and a web app."),
    ])

    # 3 Problem statement -----------------------------------------------------
    s = new_slide(prs, "Problem Statement", notes="Read the statement; stress 'any combination' and 'not a diagnosis'.")
    q = rect(s, MARGIN, Inches(1.6), W - 2 * MARGIN, Inches(1.7), TEAL)
    box_text(q, "Given a frontal chest X-ray, estimate for each of 14 findings whether it is likely present — "
             "allowing any combination — explain the evidence, and evaluate without patient-level leakage.",
             20, WHITE, align=PP_ALIGN.LEFT)
    items = [("Multi-label", "Findings co-occur; one image can be positive for several labels."),
             ("Explainable", "Grad-CAM shows which regions influenced each score."),
             ("Leakage-free", "All images of a patient stay in one split."),
             ("Honest", "Scores are experimental model outputs, never a diagnosis.")]
    for i, (h, b) in enumerate(items):
        card(s, MARGIN + i * Inches(3.05), Inches(3.75), Inches(2.85), Inches(2.6), h, b)

    # 4 Motivation ------------------------------------------------------------
    s = new_slide(prs, "Motivation", notes="Contrast common student projects with the realistic issues this one handles.")
    rows = [("Binary demos", "Pneumonia vs normal ignores co-occurring findings", "14 sigmoid outputs"),
            ("Random image splits", "Same patient in train and test inflates scores", "Patient-level splits"),
            ("Accuracy only", "Hides missed rare findings under imbalance", "AUC, PR-AUC, recall, F1"),
            ("Fixed 0.5 threshold", "Poor fit for rare labels", "Validation-tuned thresholds"),
            ("Black-box output", "No insight into model evidence", "Grad-CAM with caveats")]
    text(s, MARGIN, Inches(1.45), Inches(3.6), Inches(0.4), "Common issue", 15, MUTED, bold=True)
    text(s, Inches(4.3), Inches(1.45), Inches(5), Inches(0.4), "Why it matters", 15, MUTED, bold=True)
    text(s, Inches(9.4), Inches(1.45), Inches(3.4), Inches(0.4), "ThoraxAI approach", 15, MUTED, bold=True)
    for i, (a, b, c) in enumerate(rows):
        y = Inches(1.95) + i * Inches(0.98)
        rect(s, MARGIN, y, W - 2 * MARGIN, Inches(0.82), TEAL_TINT if i % 2 == 0 else WHITE)
        text(s, MARGIN + Inches(0.2), y, Inches(3.4), Inches(0.82), a, 16, TEAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(4.3), y, Inches(5), Inches(0.82), b, 15, INK, anchor=MSO_ANCHOR.MIDDLE)
        pill = rect(s, Inches(9.4), y + Inches(0.14), Inches(3.3), Inches(0.54), CORAL_TINT)
        box_text(pill, c, 14, INK, bold=True)

    # 5 Objectives ------------------------------------------------------------
    s = new_slide(prs, "Objectives", notes="Five objectives map directly to the code modules.")
    objs = [("Data pipeline", "Validate dataset, patient-level splits"),
            ("Model", "DenseNet121 transfer learning, 14 outputs"),
            ("Imbalance", "Weighted BCE + validation thresholds"),
            ("Evaluation", "Per-label & averaged metrics on test"),
            ("XAI + App", "Grad-CAM and Streamlit interface")]
    for i, (h, b) in enumerate(objs):
        card(s, MARGIN + i * Inches(2.45), Inches(1.7), Inches(2.25), Inches(3.4), h, b, number=i + 1)
    text(s, MARGIN, Inches(5.6), W - 2 * MARGIN, Inches(0.8),
         "Every number shown in the app or report must come from files generated by our scripts.", 16, MUTED,
         italic=True)

    # 6 Literature survey -----------------------------------------------------
    s = new_slide(prs, "Literature Survey", notes="Verify each reference before the review.")
    lit = [("Wang et al., CVPR 2017", "ChestX-ray8/14 dataset, NLP-mined labels"),
           ("Rajpurkar et al., 2017 (CheXNet)", "DenseNet121 on ChestX-ray14"),
           ("Huang et al., CVPR 2017", "DenseNet: dense connectivity"),
           ("Sandler et al., CVPR 2018", "MobileNetV2: efficient CNN"),
           ("Selvaraju et al., ICCV 2017", "Grad-CAM explanations"),
           ("Irvin et al., AAAI 2019", "CheXpert: uncertainty labels"),
           ("Zech et al., PLOS Med 2018", "Poor cross-hospital generalisation"),
           ("Seyyed-Kalantari et al., Nat Med 2021", "Under-diagnosis bias")]
    tbl = s.shapes.add_table(len(lit) + 1, 2, MARGIN, Inches(1.5), W - 2 * MARGIN, Inches(5.1)).table
    tbl.columns[0].width = Inches(5.2)
    tbl.columns[1].width = W - 2 * MARGIN - Inches(5.2)
    for c, head in enumerate(("Work", "Key contribution / relevance")):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL
        cell.text = head
        _run_style(cell.text_frame.paragraphs[0].runs[0], 15, WHITE, bold=True)
    for r, (a, b) in enumerate(lit, start=1):
        for c, val in enumerate((a, b)):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL_TINT if r % 2 else WHITE
            cell.text = val
            _run_style(cell.text_frame.paragraphs[0].runs[0], 14, INK, bold=(c == 0))

    # 7 Dataset ---------------------------------------------------------------
    s = new_slide(prs, "Dataset and the 14 Findings", notes="Quote counts only from your validation report.")
    bullets(s, MARGIN, Inches(1.5), Inches(5.2), Inches(3.5), [
        ("Source: ", "NIH ChestX-ray14 (Kaggle: nih-chest-xrays/data)"),
        ("Commonly reported: ", "112,120 images, 30,805 patients"),
        ("Labels: ", "text-mined from radiology reports (noisy)"),
        ("Splits: ", "official patient-disjoint test list preserved"),
    ], size=15)
    ph = rect(s, MARGIN, Inches(4.85), Inches(5.2), Inches(1.55), CORAL_TINT, line=CORAL)
    box_text(ph, ["PLACEHOLDER", "Insert results/dataset_figures/label_prevalence.png"], 14, INK, bold=True)
    for i, label in enumerate(config.LABELS):
        col, row = i % 2, i // 2
        chip = rect(s, Inches(6.3) + col * Inches(3.25), Inches(1.5) + row * Inches(0.7), Inches(3.05),
                    Inches(0.56), TEAL_TINT)
        box_text(chip, config.DISPLAY_NAMES[label], 15, TEAL, bold=True)

    # 8 Methodology -----------------------------------------------------------
    s = new_slide(prs, "Proposed Methodology", notes="Walk through the pipeline left to right, then the bottom row.")
    steps = ["Validate dataset", "Patient-level split", "Preprocess & augment", "Train DenseNet121",
             "Select thresholds (val)", "Evaluate (test)", "Serve + Grad-CAM"]
    bw, bh, gap = Inches(2.7), Inches(1.15), Inches(0.45)
    for i, step in enumerate(steps):
        row, k = divmod(i, 4)
        col = k if row == 0 else 3 - k  # second row runs right-to-left (snake layout)
        x = MARGIN + col * (bw + gap)
        y = Inches(1.7) + row * Inches(2.1)
        b = rect(s, x, y, bw, bh, CORAL if i == len(steps) - 1 else TEAL_MID)
        box_text(b, [f"Step {i + 1}", step], 16, WHITE, bold=True)
        if i == len(steps) - 1:
            continue
        if row == 0 and col < 3:
            arrow(s, x + bw, y + bh / 2, x + bw + gap, y + bh / 2)
        elif row == 0:
            arrow(s, x + bw / 2, y + bh, x + bw / 2, y + Inches(2.1))
        else:
            arrow(s, x, y + bh / 2, x - gap, y + bh / 2)
    text(s, MARGIN, Inches(5.75), W - 2 * MARGIN, Inches(0.8),
         "Test data is used once, at the end. Thresholds and checkpoints are chosen on validation data only.",
         15, MUTED, italic=True)

    # 9 Architecture ----------------------------------------------------------
    s = new_slide(prs, "System Architecture", notes="Offline scripts produce artifacts; the app only reads them.")
    groups = [("Offline pipeline", ["validate_dataset.py", "train_model.py", "evaluate_model.py"], TEAL_MID),
              ("Artifacts", ["models/<arch>/best_model.keras", "threshold_config.json", "results/*.json, *.png"], MUTED),
              ("Streamlit app", ["Upload & analysis", "Grad-CAM", "Dashboards"], CORAL)]
    for i, (head, lines, color) in enumerate(groups):
        x = MARGIN + i * Inches(4.15)
        rect(s, x, Inches(1.6), Inches(3.75), Inches(4.4), TEAL_TINT)
        text(s, x + Inches(0.25), Inches(1.75), Inches(3.3), Inches(0.5), head, 18, TEAL, bold=True)
        for j, line in enumerate(lines):
            b = rect(s, x + Inches(0.25), Inches(2.45) + j * Inches(1.1), Inches(3.25), Inches(0.85), color)
            box_text(b, line, 14, WHITE, bold=True)
        if i < 2:
            arrow(s, x + Inches(3.75), Inches(3.8), x + Inches(4.15), Inches(3.8))
    text(s, MARGIN, Inches(6.25), W - 2 * MARGIN, Inches(0.5),
         "Inference logic (src/predict.py, src/gradcam.py) is independent of Streamlit and unit-tested.", 15, MUTED,
         italic=True)

    # 10 Model architecture ---------------------------------------------------
    s = new_slide(prs, "Model Architecture", notes="Explain each layer; mention frozen backbone then fine-tuning.")
    layers = [("Input", "224×224×3, 0–255"), ("Normalisation", "ImageNet mean/std"),
              ("DenseNet121", "ImageNet weights → 7×7×1024"), ("Global Avg Pool", "1024 features"),
              ("Dropout 0.3", "regularisation"), ("Dense(14, sigmoid)", "14 independent scores")]
    for i, (a, b) in enumerate(layers):
        y = Inches(1.45) + i * Inches(0.92)
        fill = CORAL if i == len(layers) - 1 else (TEAL if i == 2 else TEAL_MID)
        blk = rect(s, MARGIN, y, Inches(5.6), Inches(0.75), fill)
        box_text(blk, f"{a}  —  {b}", 15, WHITE, bold=True)
    card(s, Inches(6.9), Inches(1.45), Inches(5.8), Inches(2.4), "Weighted binary cross-entropy",
         "w⁺ = N / 2P,  w⁻ = N / 2(N − P) per finding, from the training split — rare findings get larger "
         "positive weight.")
    card(s, Inches(6.9), Inches(4.1), Inches(5.8), Inches(2.4), "Two-stage training",
         "Stage 1: train head, backbone frozen (AdamW 1e-3). Stage 2 (optional): unfreeze last 40 layers, "
         "BatchNorm frozen, lr 1e-5. Early stopping on val AUC.", fill=CORAL_TINT)

    # 11 Multi-label workflow -------------------------------------------------
    s = new_slide(prs, "Multi-Label Prediction Workflow",
                  notes="The bar values are an illustrative example of the output format, not model results.")
    ex = [("Effusion", 0.81, 0.42), ("Cardiomegaly", 0.66, 0.35), ("Atelectasis", 0.30, 0.38),
          ("Infiltration", 0.22, 0.40), ("Nodule", 0.08, 0.30)]
    text(s, MARGIN, Inches(1.4), Inches(6.5), Inches(0.4), "Illustrative output format (not real results)", 14,
         MUTED, italic=True)
    for i, (name, sc, thr) in enumerate(ex):
        y = Inches(1.95) + i * Inches(0.85)
        text(s, MARGIN, y, Inches(1.9), Inches(0.55), name, 15, INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        track_x, track_w = MARGIN + Inches(2.0), Inches(4.4)
        rect(s, track_x, y + Inches(0.1), track_w, Inches(0.35), TEAL_TINT, shape=MSO_SHAPE.RECTANGLE)
        rect(s, track_x, y + Inches(0.1), Emu(int(track_w * sc)), Inches(0.35), CORAL if sc >= thr else TEAL_MID,
             shape=MSO_SHAPE.RECTANGLE)
        tx = track_x + Emu(int(track_w * thr))
        rect(s, tx, y, Inches(0.05), Inches(0.55), INK, shape=MSO_SHAPE.RECTANGLE)
    bullets(s, Inches(7.4), Inches(1.5), Inches(5.3), Inches(4.6), [
        ("Sigmoid, not softmax: ", "scores are independent and need not sum to 1."),
        ("Per-finding thresholds ", "(black marks) chosen on validation data."),
        ("Flagged ", "(coral) = score ≥ threshold; several can be flagged together."),
        ("No Finding predicted ", "= nothing crosses its threshold — a model status, not a normal result."),
    ], size=16)

    # 12 Grad-CAM -------------------------------------------------------------
    s = new_slide(prs, "Explainability with Grad-CAM", notes="Emphasise: influence, not location of disease.")
    gsteps = ["Pick a finding's score y", "Gradients ∂y/∂A on last conv maps", "Average → channel weights α",
              "ReLU(Σ α·A) → upsample → overlay"]
    for i, st in enumerate(gsteps):
        y = Inches(1.55) + i * Inches(1.08)
        badge(s, MARGIN, y + Inches(0.12), i + 1)
        text(s, MARGIN + Inches(0.7), y, Inches(5.3), Inches(0.8), st, 17, INK, anchor=MSO_ANCHOR.MIDDLE)
    ph = rect(s, Inches(7.0), Inches(1.5), Inches(5.7), Inches(3.6), CORAL_TINT, line=CORAL)
    box_text(ph, ["PLACEHOLDER", "Insert original + Grad-CAM screenshot from your trained model"], 15, INK, bold=True)
    warn = rect(s, MARGIN, Inches(5.95), W - 2 * MARGIN, Inches(0.75), TEAL)
    box_text(warn, "Grad-CAM shows what influenced the score — not where disease is, and not proof of correctness.",
             16, WHITE, bold=True)

    # 13 Results (placeholder) ------------------------------------------------
    s = new_slide(prs, "Experimental Results", notes="Fill ONLY from results/densenet121/test_metrics.json and "
                  "per_label_metrics.csv. State development vs full mode and the test-set size.")
    metric_names = ["Macro ROC-AUC", "Macro PR-AUC", "Macro F1", "Micro F1"]
    for i, m in enumerate(metric_names):
        x = MARGIN + i * Inches(3.08)
        rect(s, x, Inches(1.5), Inches(2.85), Inches(1.7), TEAL_TINT)
        text(s, x + Inches(0.2), Inches(1.6), Inches(2.5), Inches(0.9), "—", 44, CORAL, bold=True, font=HEAD_FONT)
        text(s, x + Inches(0.2), Inches(2.55), Inches(2.5), Inches(0.5), m, 15, INK, bold=True)
    for i, label in enumerate(("Per-finding metrics table", "ROC & PR curves")):
        ph = rect(s, MARGIN + i * Inches(6.15), Inches(3.5), Inches(5.95), Inches(2.4), CORAL_TINT, line=CORAL)
        box_text(ph, ["PLACEHOLDER", f"{label} from results/densenet121/"], 15, INK, bold=True)
    text(s, MARGIN, Inches(6.1), W - 2 * MARGIN, Inches(0.6),
         "Pending: training and evaluation on the real dataset. No values are shown until they are measured.",
         15, MUTED, italic=True)

    # 14 Limitations & future scope ------------------------------------------
    s = new_slide(prs, "Limitations and Future Scope", notes="Be candid; examiners value awareness of limits.")
    card(s, MARGIN, Inches(1.5), Inches(5.9), Inches(5.0), "Limitations",
         "Noisy NLP-derived labels · single hospital source · 224 px resolution · uncalibrated scores · "
         "F1-tuned (non-clinical) thresholds · possible shortcut learning · coarse Grad-CAM · no clinical validation")
    card(s, Inches(6.8), Inches(1.5), Inches(5.9), Inches(5.0), "Future scope",
         "GPU fine-tuning at higher resolution · ensembles · probability calibration · CheXpert uncertainty labels · "
         "external validation (MIMIC-CXR) · subgroup fairness audit · localisation with bounding boxes · DICOM",
         fill=CORAL_TINT)

    # 15 Conclusion & references ---------------------------------------------
    s = new_slide(prs, "Conclusion", dark=True, notes="Summarise and invite questions.")
    bullets(s, MARGIN, Inches(1.5), Inches(6.2), Inches(4.5), [
        "End-to-end, tested pipeline for 14-finding multi-label classification",
        "Leakage-free patient-level evaluation with imbalance-aware metrics",
        "Grad-CAM explanations with explicit caveats",
        "Application that never fabricates predictions or metrics",
    ], size=17, color=WHITE)
    refs = rect(s, Inches(7.2), Inches(1.5), Inches(5.5), Inches(4.9), TEAL_MID)
    box_text(refs, ["Key references",
                    "Wang et al., CVPR 2017 — ChestX-ray14",
                    "Rajpurkar et al., 2017 — CheXNet",
                    "Huang et al., CVPR 2017 — DenseNet",
                    "Selvaraju et al., ICCV 2017 — Grad-CAM",
                    "Irvin et al., AAAI 2019 — CheXpert",
                    "Zech et al., PLOS Med 2018 — generalisation"], 17, WHITE, bold=True, align=PP_ALIGN.LEFT)
    text(s, MARGIN, Inches(6.4), Inches(6), Inches(0.6), "Thank you — questions?", 22, CORAL, bold=True,
         font=HEAD_FONT)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print("Saved", build())
